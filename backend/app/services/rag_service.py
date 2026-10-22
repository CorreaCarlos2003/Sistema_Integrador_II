"""
Motor de RAG (Retrieval-Augmented Generation) del proyecto.

Reimplementa, para uso desde la API, la lógica prototipada en el notebook
`RAG/Base_Chroma.ipynb`: carga de .docx, chunking y búsqueda semántica con
Chroma. Cada proyecto tiene su propia colección de Chroma, aislada de las
demás, persistida en disco bajo `RAG_STORAGE_DIR`.

Por ahora solo se implementa retrieval (no hay generación de respuesta con
un LLM): el chat devuelve los fragmentos más relevantes para la pregunta.
"""

import os
import shutil

from docx import Document as WordDocument
from docx.table import Table
from docx.text.paragraph import Paragraph
from openpyxl import load_workbook
from pptx import Presentation
from pypdf import PdfReader
from langchain_core.documents import Document
from langchain_text_splitters import RecursiveCharacterTextSplitter
from langchain_huggingface import HuggingFaceEmbeddings
from langchain_chroma import Chroma

from app.core.config import (
    RAG_CHUNK_OVERLAP,
    RAG_CHUNK_SIZE,
    RAG_EMBEDDING_MODEL,
    RAG_STORAGE_DIR,
)

_embeddings = None


def _obtener_embeddings() -> HuggingFaceEmbeddings:
    """Carga el modelo de embeddings una sola vez (es costoso de inicializar)."""
    global _embeddings
    if _embeddings is None:
        _embeddings = HuggingFaceEmbeddings(model_name=RAG_EMBEDDING_MODEL)
    return _embeddings


def _directorio_proyecto(proyecto_id: int) -> str:
    ruta = os.path.join(RAG_STORAGE_DIR, f"proyecto_{proyecto_id}")
    os.makedirs(os.path.join(ruta, "documentos"), exist_ok=True)
    return ruta


def ruta_documentos(proyecto_id: int) -> str:
    return os.path.join(_directorio_proyecto(proyecto_id), "documentos")


def obtener_vectorstore(proyecto_id: int) -> Chroma:
    directorio = os.path.join(_directorio_proyecto(proyecto_id), "chroma_db")
    return Chroma(
        embedding_function=_obtener_embeddings(),
        persist_directory=directorio,
        collection_name=f"proyecto_{proyecto_id}",
    )


def extraer_bloques_docx(ruta: str) -> list[tuple[str, str]]:
    """Extrae párrafos y tablas en el orden en que aparecen en el DOCX."""
    docx = WordDocument(ruta)
    bloques = []
    for elemento in docx.element.body.iterchildren():
        if elemento.tag.endswith("}p"):
            parrafo = Paragraph(elemento, docx)
            texto = parrafo.text.strip()
            if texto:
                nombre_estilo = parrafo.style.name if parrafo.style is not None else "Normal"
                bloques.append((texto, nombre_estilo))
        elif elemento.tag.endswith("}tbl"):
            tabla = Table(elemento, docx)
            for fila in tabla.rows:
                texto = " | ".join(celda.text.strip().replace("\n", " / ") for celda in fila.cells)
                if texto.strip(" |"):
                    bloques.append((texto, "Tabla"))
    return bloques


def _cargar_docx(ruta: str, nombre: str | None = None) -> Document:
    contenido = "\n".join(texto for texto, _ in extraer_bloques_docx(ruta))

    return Document(
        page_content=contenido,
        metadata={
            "source": ruta,
            "nombre": nombre or os.path.basename(ruta),
            "tipo": "docx",
        },
    )


def _cargar_pdf(ruta: str, nombre: str | None = None) -> Document:
    lector = PdfReader(ruta)
    paginas = [(pagina.extract_text() or "").strip() for pagina in lector.pages]
    contenido = "\n".join(texto for texto in paginas if texto)

    return Document(
        page_content=contenido,
        metadata={
            "source": ruta,
            "nombre": nombre or os.path.basename(ruta),
            "tipo": "pdf",
        },
    )


def _cargar_excel(ruta: str, nombre: str | None = None) -> Document:
    libro = load_workbook(ruta, data_only=True, read_only=True)
    bloques = []
    for hoja in libro.worksheets:
        for fila in hoja.iter_rows(values_only=True):
            celdas = [str(celda).strip() for celda in fila if celda is not None and str(celda).strip()]
            if celdas:
                bloques.append(" | ".join(celdas))
    contenido = "\n".join(bloques)

    return Document(
        page_content=contenido,
        metadata={
            "source": ruta,
            "nombre": nombre or os.path.basename(ruta),
            "tipo": "xlsx",
        },
    )


def _cargar_pptx(ruta: str, nombre: str | None = None) -> Document:
    presentacion = Presentation(ruta)
    bloques = []
    for indice, diapositiva in enumerate(presentacion.slides, start=1):
        textos = []
        for forma in diapositiva.shapes:
            if forma.has_text_frame:
                texto = forma.text_frame.text.strip()
                if texto:
                    textos.append(texto)
            if forma.has_table:
                for fila in forma.table.rows:
                    texto_fila = " | ".join(celda.text.strip() for celda in fila.cells)
                    if texto_fila.strip(" |"):
                        textos.append(texto_fila)
        if textos:
            bloques.append(f"[Diapositiva {indice}]\n" + "\n".join(textos))
    contenido = "\n\n".join(bloques)

    return Document(
        page_content=contenido,
        metadata={
            "source": ruta,
            "nombre": nombre or os.path.basename(ruta),
            "tipo": "pptx",
        },
    )


def _cargar_texto_plano(ruta: str, nombre: str | None = None, tipo: str = "txt") -> Document:
    with open(ruta, "r", encoding="utf-8", errors="replace") as archivo:
        contenido = archivo.read()

    return Document(
        page_content=contenido,
        metadata={
            "source": ruta,
            "nombre": nombre or os.path.basename(ruta),
            "tipo": tipo,
        },
    )


_CARGADORES_POR_EXTENSION = {
    ".docx": _cargar_docx,
    ".pdf": _cargar_pdf,
    ".xlsx": _cargar_excel,
    ".pptx": _cargar_pptx,
    ".txt": lambda ruta, nombre=None: _cargar_texto_plano(ruta, nombre, "txt"),
    ".md": lambda ruta, nombre=None: _cargar_texto_plano(ruta, nombre, "md"),
}


def _cargar_documento(ruta: str, nombre: str | None = None) -> Document:
    extension = os.path.splitext(nombre or ruta)[1].lower()
    cargador = _CARGADORES_POR_EXTENSION.get(extension)
    if cargador is None:
        raise ValueError(f"Formato de archivo no soportado: {extension}")
    return cargador(ruta, nombre)


def indexar_documento(proyecto_id: int, ruta_archivo: str, nombre: str | None = None) -> int:
    """Parte un documento en chunks, los embebe y los guarda en la colección del proyecto.

    Soporta .docx, .pdf, .xlsx, .pptx, .txt y .md.
    Devuelve la cantidad de chunks indexados.
    """
    documento = _cargar_documento(ruta_archivo, nombre)

    splitter = RecursiveCharacterTextSplitter(
        chunk_size=RAG_CHUNK_SIZE,
        chunk_overlap=RAG_CHUNK_OVERLAP,
    )
    chunks = splitter.split_documents([documento])

    for i, chunk in enumerate(chunks):
        chunk.metadata["chunk_id"] = i
        chunk.metadata["proyecto_id"] = proyecto_id

    if not chunks:
        return 0

    vectorstore = obtener_vectorstore(proyecto_id)
    vectorstore.add_documents(chunks)

    return len(chunks)


def eliminar_documento_indexado(proyecto_id: int, ruta_archivo: str) -> None:
    """Elimina de la colección del proyecto todos los chunks de ese archivo."""
    vectorstore = obtener_vectorstore(proyecto_id)
    vectorstore.delete(where={"source": ruta_archivo})


def eliminar_proyecto_indexado(proyecto_id: int) -> None:
    """Borra del disco los documentos y la colección vectorial del proyecto."""
    ruta = os.path.join(RAG_STORAGE_DIR, f"proyecto_{proyecto_id}")
    if os.path.isdir(ruta):
        shutil.rmtree(ruta)


def buscar(
    proyecto_id: int, pregunta: str, k: int = 5, rutas_archivo: list[str] | None = None,
) -> list[dict]:
    vectorstore = obtener_vectorstore(proyecto_id)
    if rutas_archivo is not None:
        if not rutas_archivo:
            return []
        filtro = (
            {"source": rutas_archivo[0]}
            if len(rutas_archivo) == 1
            else {"source": {"$in": rutas_archivo}}
        )
        resultados = vectorstore.similarity_search_with_score(pregunta, k=k, filter=filtro)
    else:
        resultados = vectorstore.similarity_search_with_score(pregunta, k=k)

    return [
        {
            "documento": doc.metadata.get("nombre", ""),
            "chunk_id": doc.metadata.get("chunk_id", -1),
            "contenido": doc.page_content,
            "score": float(score),
            "ruta_archivo": doc.metadata.get("source", ""),
            "pagina": doc.metadata.get("pagina"),
            "seccion": doc.metadata.get("seccion"),
        }
        for doc, score in resultados
    ]
